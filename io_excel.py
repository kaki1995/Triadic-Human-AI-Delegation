# triadic_sim/io_excel.py
from __future__ import annotations

from typing import Dict, List
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


def safe_sheet_to_df(ws: Worksheet) -> pd.DataFrame:
    """Read a worksheet where the first row is a header."""
    rows = list(ws.values)
    if not rows:
        return pd.DataFrame()
    header = [str(x).strip() if x is not None else "" for x in rows[0]]
    data = rows[1:]
    df = pd.DataFrame(data, columns=header)
    df = df.loc[:, [c for c in df.columns if c and c != "None"]]
    return df


def write_df_to_sheet(ws: Worksheet, df: pd.DataFrame) -> None:
    """Overwrite worksheet with DataFrame content."""
    ws.delete_rows(1, ws.max_row if ws.max_row else 1)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))


def ensure_columns(df: pd.DataFrame, cols: List[str]) -> pd.DataFrame:
    """Ensure df has all columns, add missing columns as NA, return in same order."""
    for c in cols:
        if c not in df.columns:
            df[c] = pd.NA
    return df[cols]


def write_to_schema_workbook(
    input_schema_xlsx: str,
    output_xlsx: str,
    dfs: Dict[str, pd.DataFrame],
    sheet_map: Dict[str, str],
) -> None:
    """
    Write DataFrames into an existing schema workbook.

    Behavior:
    - If sheet exists: align df columns to existing sheet header columns and overwrite
    - If sheet missing: create it and write df as-is
    """
    wb = load_workbook(input_schema_xlsx)

    for df_key, sheet_name in sheet_map.items():
        df = dfs[df_key].copy()

        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            existing = safe_sheet_to_df(ws)
            if not existing.empty:
                df = ensure_columns(df, list(existing.columns))
            write_df_to_sheet(ws, df)
        else:
            ws = wb.create_sheet(sheet_name)
            write_df_to_sheet(ws, df)

    wb.save(output_xlsx)
